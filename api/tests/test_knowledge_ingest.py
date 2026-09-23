"""Tests for services/knowledge_ingest.py's AI ingestion pass, ported over
from a bug found on the sibling single-tenant project (shofar-cloud's
dev-whatsapp/api/shofar_automation/services/knowledge_ingest.py): a
truncated JSON response from the ingestion model (most commonly caused by
a source document producing more FAQ content than max_tokens allows)
raised a raw, unhandled pydantic.ValidationError that propagated past
web/knowledge_router.py's `except IngestError` handling and 500'd the
request instead of showing the staff member a clean message.

No async test plugin (pytest-asyncio/anyio) is installed in this repo -
driven via asyncio.run() in plain sync test functions, same as
test_ai_reply_safety_nets.py. clients/storage calls are monkeypatched
directly on knowledge_ingest's own module references (it does
`from autosend import clients, storage`), so nothing here touches a real
DB or makes a real Anthropic API call."""
import asyncio

import pytest

from autosend.services import knowledge_ingest


class _TruncatingMessages:
    """Stands in for AsyncAnthropic().messages - .parse() raises a real
    pydantic.ValidationError, the same way the SDK's own JSON-mode parsing
    fails when the model's output got cut off mid-string before it formed
    valid JSON matching the requested output_format."""

    async def parse(self, **kwargs):
        knowledge_ingest._FAQPairs.model_validate({"pairs": "not-a-list"})


class _TruncatingAnthropicClient:
    def __init__(self):
        self.messages = _TruncatingMessages()


def test_generate_faq_pairs_converts_truncated_json_to_ingest_error(monkeypatch):
    # The Anthropic key itself is the shared platform-wide credential
    # (clients.get_anthropic_client(), not a per-pipeline copy) - only the
    # model is ingestion's own setting.
    monkeypatch.setattr(
        knowledge_ingest.storage, "get_ai_ingestion_settings",
        lambda: {"model": "claude-x"},
    )
    monkeypatch.setattr(
        knowledge_ingest.clients, "get_anthropic_client", lambda: _TruncatingAnthropicClient(),
    )

    def _fail_if_called(**kwargs):
        raise AssertionError("record_ai_ingestion should not run once parsing has failed")

    monkeypatch.setattr(knowledge_ingest.storage, "record_ai_ingestion", _fail_if_called)

    with pytest.raises(knowledge_ingest.IngestError) as exc_info:
        asyncio.run(
            knowledge_ingest._generate_faq_pairs(
                "some long source text", "Some Document",
                org_id=1, unit_id=1, source_type="pdf", source_ref="doc.pdf",
            )
        )

    assert "more FAQ content than a single pass can handle" in str(exc_info.value)


def test_generate_faq_pairs_requires_a_model_even_though_key_is_shared(monkeypatch):
    # No model configured under the Ingestion tab yet - this must fail
    # before ever touching the (possibly unconfigured) shared Anthropic
    # credential, since the two are independent settings now.
    monkeypatch.setattr(knowledge_ingest.storage, "get_ai_ingestion_settings", lambda: {"model": None})

    def _fail_if_called():
        raise AssertionError("get_anthropic_client should not run without a model configured")

    monkeypatch.setattr(knowledge_ingest.clients, "get_anthropic_client", _fail_if_called)

    with pytest.raises(knowledge_ingest.IngestError) as exc_info:
        asyncio.run(
            knowledge_ingest._generate_faq_pairs(
                "some text", "Some Document", org_id=1, unit_id=1, source_type="pdf", source_ref="doc.pdf",
            )
        )

    assert "Ingestion tab" in str(exc_info.value)


def test_generate_faq_pairs_surfaces_missing_shared_key_as_ingest_error(monkeypatch):
    # get_anthropic_client() raises a plain ValueError when ai_credentials
    # isn't configured yet (see clients.py) - this must become a clean
    # IngestError for the router, not an unhandled 500.
    monkeypatch.setattr(knowledge_ingest.storage, "get_ai_ingestion_settings", lambda: {"model": "claude-x"})

    def _raise_unconfigured():
        raise ValueError(
            "AI credentials aren't configured yet - a superadmin needs to add an "
            "Anthropic API key under AI Credentials first."
        )

    monkeypatch.setattr(knowledge_ingest.clients, "get_anthropic_client", _raise_unconfigured)

    with pytest.raises(knowledge_ingest.IngestError) as exc_info:
        asyncio.run(
            knowledge_ingest._generate_faq_pairs(
                "some text", "Some Document", org_id=1, unit_id=1, source_type="pdf", source_ref="doc.pdf",
            )
        )

    assert "AI credentials aren't configured yet" in str(exc_info.value)


# --- ingest_upload: multi-format text extraction ---------------------------
# Built in memory with the same libraries the extractors read with, so these
# exercise the real parsing path without shipping binary fixtures.

def _docx_bytes() -> bytes:
    import io
    import docx
    document = docx.Document()
    document.add_paragraph("Service times")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Sunday"
    table.rows[0].cells[1].text = "09:00"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    import io
    import pptx
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Youth camp"
    slide.notes_slide.notes_text_frame.text = "Contact Sipho on 082 000 0000"
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import io
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.active.title = "Prices"
    workbook.active.append(["Ticket", "R150"])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_extract_docx_keeps_paragraphs_and_table_rows():
    text = knowledge_ingest._extract_docx_text(_docx_bytes())
    assert "Service times" in text
    assert "Sunday | 09:00" in text


def test_extract_pptx_includes_slide_text_and_speaker_notes():
    text = knowledge_ingest._extract_pptx_text(_pptx_bytes())
    assert "Slide 1:" in text
    assert "Youth camp" in text
    assert "Notes: Contact Sipho on 082 000 0000" in text


def test_extract_xlsx_labels_sheets_and_joins_cells():
    text = knowledge_ingest._extract_xlsx_text(_xlsx_bytes())
    assert "Sheet: Prices" in text
    assert "Ticket | R150" in text


def test_decode_text_falls_back_to_cp1252():
    assert knowledge_ingest._decode_text("café".encode("cp1252")) == "café"
    assert knowledge_ingest._decode_text("﻿hello".encode("utf-8")) == "hello"


def test_extract_html_strips_scripts():
    text = knowledge_ingest._extract_html_text(b"<html><script>evil()</script><p>Hello there</p></html>")
    assert "Hello there" in text
    assert "evil" not in text


@pytest.mark.parametrize("filename,expected", [
    ("old.doc", ".docx"), ("old.ppt", ".pptx"), ("old.xls", ".xlsx"), ("image.png", "Unsupported"),
])
def test_ingest_upload_rejects_legacy_and_unknown_formats_before_ai_pass(monkeypatch, filename, expected):
    async def _fail_if_called(*args, **kwargs):
        raise AssertionError("_generate_faq_pairs must not run for a rejected file type")

    monkeypatch.setattr(knowledge_ingest, "_generate_faq_pairs", _fail_if_called)
    with pytest.raises(knowledge_ingest.IngestError) as exc_info:
        asyncio.run(knowledge_ingest.ingest_upload(1, None, filename, b"x" * 100))
    assert expected in str(exc_info.value)


@pytest.mark.parametrize("filename,source_type", [("guide.docx", "file"), ("Guide.PDF", "pdf")])
def test_ingest_upload_source_type_keeps_pdf_distinct(monkeypatch, filename, source_type):
    """PDFs keep source_type 'pdf' so a re-upload still replaces chunks
    ingested before other formats existed; everything else is 'file'."""
    captured = {}

    async def _fake_pairs(text, title, **kwargs):
        captured["pairs_source_type"] = kwargs["source_type"]
        return [{"title": "Q", "content": "A"}]

    def _fake_replace(org_id, unit_id, st, source_ref, pairs, *, document_title=None):
        captured["replace_source_type"] = st
        return [1]

    monkeypatch.setattr(knowledge_ingest, "_generate_faq_pairs", _fake_pairs)
    monkeypatch.setattr(knowledge_ingest, "_UPLOAD_EXTRACTORS", {
        **knowledge_ingest._UPLOAD_EXTRACTORS, ".pdf": lambda b: "x" * 100, ".docx": lambda b: "x" * 100,
    })
    monkeypatch.setattr(knowledge_ingest.storage, "get_knowledge_base_source_document_title", lambda *a: None)
    monkeypatch.setattr(knowledge_ingest.storage, "replace_knowledge_base_source_entries", _fake_replace)

    assert asyncio.run(knowledge_ingest.ingest_upload(1, None, filename, b"irrelevant")) == [1]
    assert captured == {"pairs_source_type": source_type, "replace_source_type": source_type}
