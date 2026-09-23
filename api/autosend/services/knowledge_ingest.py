"""Knowledge Base ingestion for the AI Assistant module.

Three input shapes, all funnelled through the same AI "FAQ-ification"
pass (_generate_faq_pairs) except manual entry (already clean, no AI pass
needed):
  - save_manual_qa()  - a staff-typed Q&A pair, stored verbatim.
  - scrape_url()      - httpx GET + BeautifulSoup cleanup -> FAQ pass.
  - ingest_upload()   - text extraction from an uploaded document (PDF,
                        Word, PowerPoint, Excel, text/Markdown/CSV or
                        HTML, see _UPLOAD_EXTRACTORS) -> FAQ pass.

Not implemented here (scoped out of this pass): a headless-browser
fallback for JS-rendered pages where the static httpx fetch yields too
little text. That needs a Playwright/Chromium install with OS-level
dependencies baked into the Docker image, which is a bigger, separate
change to api/Dockerfile - scrape_url() below only does the static fetch.
"""
from __future__ import annotations

import io

import anyio
import docx
import httpx
import openpyxl
import pptx
import pydantic
from bs4 import BeautifulSoup
from docx.table import Table as DocxTable
from pydantic import BaseModel
from pypdf import PdfReader

from autosend import clients, storage
from autosend.utils.logging import get_logger

logger = get_logger(__name__)

_USER_AGENT = "Mozilla/5.0 (compatible; KryxBot/1.0; +https://kryx.co.za)"
_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_MAX_PDF_PAGES = 200
_MAX_INGEST_CHARS = 60_000
_MIN_TEXT_CHARS = 50

_FAQ_SYSTEM_PROMPT = """You turn raw text scraped from a website or uploaded document into a set of clear,
self-contained FAQ-style question/answer pairs for a WhatsApp AI assistant's
knowledge base.

Rules:
- Each answer must be understandable on its own, with no reference to
  "the page above" or similar context the reader won't have.
- Prefer several short, specific pairs over one long generic pair.
- Only include information actually present in the text - never invent
  facts, dates, prices, or contact details.
- Skip navigation boilerplate, cookie notices, and anything that isn't a
  real answerable fact about the organisation.
- Write answers in a warm, conversational tone suitable for a WhatsApp reply."""


class IngestError(Exception):
    """Raised for a clean, user-facing ingestion failure (bad input,
    AI ingestion settings not configured, etc) - callers should surface
    str(exc) directly rather than a generic 500."""


class _FAQPair(BaseModel):
    question: str
    answer: str


class _FAQPairs(BaseModel):
    pairs: list[_FAQPair]


def _clean_html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "noscript", "form", "iframe"]):
        tag.decompose()
    for el in soup.find_all(attrs={"aria-hidden": "true"}):
        el.decompose()
    for el in soup.select('[style*="display:none"], [style*="display: none"]'):
        el.decompose()

    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


async def _generate_faq_pairs(
    text: str, title: str, *, org_id: int, unit_id: int | None, source_type: str, source_ref: str | None,
) -> list[dict]:
    settings = storage.get_ai_ingestion_settings()
    model = settings.get("model") if settings else None
    if not model:
        raise IngestError(
            "AI ingestion settings aren't configured yet - a superadmin needs to select a "
            "model under the AI Credentials page's Ingestion tab first."
        )

    # The Anthropic key itself is the one shared platform-wide credential
    # (ai_credentials.api_key, see clients.get_anthropic_client()'s
    # docstring) - not a second copy configured here.
    try:
        client = clients.get_anthropic_client()
    except ValueError as exc:
        raise IngestError(str(exc)) from exc
    try:
        response = await client.messages.parse(
            model=model,
            # A large source document can legitimately produce a long list
            # of FAQ pairs; too low a ceiling here truncates the model's
            # JSON output mid-string, which surfaces as the
            # pydantic.ValidationError caught below. Kept well under the
            # Anthropic SDK's own ~21333 "expected_time > 10 minutes" cutoff
            # for a single non-streaming call (an uncatchable ValueError,
            # not an anthropic.APIError) - 32000 previously took down every
            # PDF/URL ingestion in production on the sibling shofar-cloud
            # project.
            max_tokens=16000,
            system=_FAQ_SYSTEM_PROMPT,
            messages=[{
                "role": "user",
                "content": f"Document title: {title}\n\nSource text:\n{text[:_MAX_INGEST_CHARS]}",
            }],
            output_format=_FAQPairs,
        )
    except pydantic.ValidationError as exc:
        # The model's JSON output didn't parse - most commonly cut off
        # mid-string because the document produced more FAQ content than
        # max_tokens above allows. Without this, the raw ValidationError
        # propagated uncaught past the router's `except IngestError`
        # handling and 500'd the request instead of showing the staff
        # member a clean message.
        raise IngestError(
            "The AI ingestion pass produced an incomplete response, likely because this "
            "document generated more FAQ content than a single pass can handle. Try "
            "splitting it into smaller documents."
        ) from exc

    output = response.parsed_output
    if output is None or not output.pairs:
        raise IngestError("The AI couldn't extract any usable FAQ pairs from this content.")

    storage.record_ai_ingestion(
        org_id=org_id, unit_id=unit_id, source_type=source_type, source_ref=source_ref,
        prompt_tokens=response.usage.input_tokens, completion_tokens=response.usage.output_tokens, model=model,
    )
    return [{"title": p.question, "content": p.answer} for p in output.pairs]


def save_manual_qa(org_id: int, unit_id: int | None, title: str, content: str) -> int:
    """A staff-typed Q&A pair, stored verbatim - already clean, so this
    skips the AI FAQ-ification pass entirely."""
    return storage.create_knowledge_base_entry(org_id, unit_id, title, content, source_type="manual")


async def scrape_url(org_id: int, unit_id: int | None, url: str, title: str | None = None) -> list[int]:
    """`title` is an optional staff-supplied document title, given to the
    AI extraction pass as context and stamped onto the resulting entries
    (storage.get_knowledge_base_source_document_title). Left blank, a
    re-scrape of a URL that already has one (whether previously set
    manually or scraped) keeps that title rather than reverting to
    whatever the page's <title> happens to be right now; only a URL with
    no prior title at all falls back to the page's <title>, then the URL
    itself."""
    try:
        async with httpx.AsyncClient(timeout=30, follow_redirects=True, headers={"User-Agent": _USER_AGENT}) as client:
            response = await client.get(url)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise IngestError(f"Couldn't fetch {url}: {exc}") from exc

    text = _clean_html_to_text(response.text)
    if len(text) < _MIN_TEXT_CHARS:
        raise IngestError(
            "This page had no usable text content to ingest (it may be a JS-rendered "
            "page this scraper can't read - try pasting the content manually instead)."
        )

    resolved_title = (title or "").strip()
    if not resolved_title:
        resolved_title = storage.get_knowledge_base_source_document_title(org_id, unit_id, "url", url) or ""
    if not resolved_title:
        page_soup = BeautifulSoup(response.text, "html.parser")
        resolved_title = page_soup.title.get_text(strip=True) if page_soup.title else ""
    resolved_title = resolved_title or url

    pairs = await _generate_faq_pairs(
        text, resolved_title, org_id=org_id, unit_id=unit_id, source_type="url", source_ref=url,
    )
    return storage.replace_knowledge_base_source_entries(
        org_id, unit_id, "url", url, pairs, document_title=resolved_title,
    )


def _extract_pdf_text(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as exc:
        raise IngestError(f"Couldn't read this PDF: {exc}") from exc
    if len(reader.pages) > _MAX_PDF_PAGES:
        raise IngestError(f"PDF has too many pages (max {_MAX_PDF_PAGES}).")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _table_rows_text(rows) -> str:
    """One line per row, cells joined with " | " - keeps a schedule/price
    table's row structure readable to the AI ingestion pass instead of
    flattening every cell into one run-on paragraph."""
    lines = []
    for row in rows:
        cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _docx_paragraph_text(paragraph) -> str:
    # paragraph.text alone keeps a hyperlink's display text but drops its
    # URL, so a document's registration/giving links would never be
    # quotable by the AI reply - inline the address after the link text.
    parts = []
    for item in paragraph.iter_inner_content():
        text = item.text or ""
        address = getattr(item, "address", None)
        parts.append(f"{text} ({address})" if address and address != text else text)
    return "".join(parts).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise IngestError(f"Couldn't read this file as a Word document: {exc}") from exc
    # iter_inner_content keeps paragraphs and tables in document order, so
    # a table stays next to the heading that introduces it.
    blocks = []
    for block in document.iter_inner_content():
        if isinstance(block, DocxTable):
            text = _table_rows_text([cell.text for cell in row.cells] for row in block.rows)
        else:
            text = _docx_paragraph_text(block)
        if text:
            blocks.append(text)
    return "\n".join(blocks)


def _extract_pptx_text(file_bytes: bytes) -> str:
    try:
        presentation = pptx.Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise IngestError(f"Couldn't read this file as a PowerPoint presentation: {exc}") from exc
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.append(_table_rows_text([cell.text for cell in row.cells] for row in shape.table.rows))
        # Speaker notes often carry the actual detail (times, contacts)
        # behind a slide's bullet-point headline.
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        parts = [p for p in parts if p]
        if parts:
            slides.append(f"Slide {number}:\n" + "\n".join(parts))
    return "\n\n".join(slides)


def _extract_xlsx_text(file_bytes: bytes) -> str:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise IngestError(f"Couldn't read this file as an Excel workbook: {exc}") from exc
    try:
        sheets = []
        for sheet in workbook.worksheets:
            text = _table_rows_text(sheet.iter_rows(values_only=True))
            if text:
                sheets.append(f"Sheet: {sheet.title}\n{text}")
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def _decode_text(file_bytes: bytes) -> str:
    # utf-8-sig strips the BOM Notepad/Excel put on "UTF-8" exports; cp1252
    # is the fallback for older Windows-saved .txt/.csv files, which would
    # otherwise fail outright on a single smart quote or accented name.
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return file_bytes.decode("cp1252", errors="replace")


def _extract_html_text(file_bytes: bytes) -> str:
    return _clean_html_to_text(_decode_text(file_bytes))


# Extension -> text extractor for ingest_upload. Keyed by extension rather
# than the browser-supplied content type, which is unreliable across
# browsers/OSes for anything other than PDF (a .md file often arrives as
# application/octet-stream).
_UPLOAD_EXTRACTORS = {
    ".pdf": _extract_pdf_text,
    ".docx": _extract_docx_text,
    ".pptx": _extract_pptx_text,
    ".xlsx": _extract_xlsx_text,
    ".txt": _decode_text,
    ".md": _decode_text,
    ".csv": _decode_text,
    ".html": _extract_html_text,
    ".htm": _extract_html_text,
}
SUPPORTED_UPLOAD_EXTENSIONS = tuple(_UPLOAD_EXTRACTORS)

# The legacy binary Office formats need a separate converter (e.g.
# LibreOffice) to read at all, which isn't worth bundling into the image
# when re-saving as the modern format is a one-click fix for staff.
_LEGACY_OFFICE_HINTS = {
    ".doc": "Word 97-2003 (.doc) files aren't supported. Open it in Word and save it as .docx first.",
    ".ppt": "PowerPoint 97-2003 (.ppt) files aren't supported. Save it as .pptx first.",
    ".xls": "Excel 97-2003 (.xls) files aren't supported. Save it as .xlsx first.",
}


def upload_source_type(filename: str) -> str:
    """PDFs keep their original 'pdf' source_type so re-uploading a PDF
    ingested before other formats were supported still replaces its
    existing chunks (source_type is part of replace_source_entries'
    dedup key); every other format shares 'file', with the extension on
    source_ref (the filename) telling them apart for display."""
    return "pdf" if filename.lower().endswith(".pdf") else "file"


async def ingest_upload(
    org_id: int, unit_id: int | None, filename: str, file_bytes: bytes, title: str | None = None,
) -> list[int]:
    """Extracts text from an uploaded document (see _UPLOAD_EXTRACTORS for
    the supported formats) and runs it through the FAQ pass. `title` is an
    optional staff-supplied document title - same "left blank keeps the
    prior title" re-ingest behaviour as scrape_url's, see its docstring.
    Only a filename with no prior title at all falls back to the raw
    filename itself."""
    if len(file_bytes) > _MAX_UPLOAD_BYTES:
        raise IngestError(f"File is too large (max {_MAX_UPLOAD_BYTES // (1024 * 1024)}MB).")
    extension = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if extension in _LEGACY_OFFICE_HINTS:
        raise IngestError(_LEGACY_OFFICE_HINTS[extension])
    extractor = _UPLOAD_EXTRACTORS.get(extension)
    if extractor is None:
        raise IngestError("Unsupported file type. Supported types: " + ", ".join(SUPPORTED_UPLOAD_EXTENSIONS))

    # python-docx/python-pptx/openpyxl/pypdf parsing is synchronous and can
    # take a while on a large file, so keep it off the event loop.
    text = await anyio.to_thread.run_sync(extractor, file_bytes)
    if len(text.strip()) < _MIN_TEXT_CHARS:
        raise IngestError("Couldn't extract any usable text from this file (it may be scanned images).")

    source_type = upload_source_type(filename)
    resolved_title = (title or "").strip()
    if not resolved_title:
        resolved_title = storage.get_knowledge_base_source_document_title(org_id, unit_id, source_type, filename) or ""
    resolved_title = resolved_title or filename

    pairs = await _generate_faq_pairs(
        text, resolved_title, org_id=org_id, unit_id=unit_id, source_type=source_type, source_ref=filename,
    )
    return storage.replace_knowledge_base_source_entries(
        org_id, unit_id, source_type, filename, pairs, document_title=resolved_title,
    )
